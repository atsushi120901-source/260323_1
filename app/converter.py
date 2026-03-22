import io
from typing import List, Optional, Tuple
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Cm, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SLIDE_SIZES = {
    "16:9": (Inches(13.33), Inches(7.5)),
    "4:3": (Inches(10), Inches(7.5)),
    "A4": (Cm(29.7), Cm(21.0)),
}


def _calc_contain(
    img_w: int, img_h: int, slide_w: Emu, slide_h: Emu, margin: float
) -> Tuple[Emu, Emu, Emu, Emu]:
    """Return (left, top, width, height) keeping aspect ratio within slide."""
    avail_w = slide_w * (1 - 2 * margin)
    avail_h = slide_h * (1 - 2 * margin)
    img_ratio = img_w / img_h
    avail_ratio = avail_w / avail_h
    if img_ratio > avail_ratio:
        w = avail_w
        h = int(w / img_ratio)
    else:
        h = avail_h
        w = int(h * img_ratio)
    left = int((slide_w - w) / 2)
    top = int((slide_h - h) / 2)
    return left, top, w, h


def _calc_fill(
    img_w: int, img_h: int, slide_w: Emu, slide_h: Emu
) -> Tuple[Emu, Emu, Emu, Emu]:
    """Return (left, top, width, height) filling the slide (may crop)."""
    img_ratio = img_w / img_h
    slide_ratio = slide_w / slide_h
    if img_ratio > slide_ratio:
        h = slide_h
        w = int(h * img_ratio)
    else:
        w = slide_w
        h = int(w / img_ratio)
    left = int((slide_w - w) / 2)
    top = int((slide_h - h) / 2)
    return left, top, w, h


def convert_images_to_pptx(
    image_data_list: List[Tuple[bytes, str]],
    slide_size: str = "16:9",
    layout: str = "contain",
    title_prefix: Optional[str] = None,
    margin: float = 0.05,
) -> io.BytesIO:
    """
    Convert a list of images to a PPTX file.

    Args:
        image_data_list: List of (image_bytes, filename) tuples.
        slide_size: One of '16:9', '4:3', 'A4'.
        layout: One of 'fit' (same as contain, no margin), 'contain', 'fill'.
        title_prefix: Optional title text to show on each slide.
        margin: Fraction of slide dimension to use as margin (0.0–0.2).

    Returns:
        BytesIO containing the PPTX file.
    """
    prs = Presentation()

    slide_w, slide_h = SLIDE_SIZES.get(slide_size, SLIDE_SIZES["16:9"])
    prs.slide_width = slide_w
    prs.slide_height = slide_h

    # Use blank slide layout
    blank_layout = prs.slide_layouts[6]

    title_height = Inches(0.5) if title_prefix is not None else 0
    # Reserve bottom area for title
    image_area_h = slide_h - title_height

    for idx, (img_bytes, filename) in enumerate(image_data_list):
        slide = prs.slides.add_slide(blank_layout)

        # Set white background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        img = Image.open(io.BytesIO(img_bytes))
        img_w, img_h = img.size

        if layout == "fit":
            left, top, w, h = _calc_contain(img_w, img_h, slide_w, image_area_h, 0.0)
        elif layout == "fill":
            left, top, w, h = _calc_fill(img_w, img_h, slide_w, image_area_h)
        else:  # contain (default)
            left, top, w, h = _calc_contain(img_w, img_h, slide_w, image_area_h, margin)

        img_stream = io.BytesIO(img_bytes)
        slide.shapes.add_picture(img_stream, left, top, w, h)

        if title_prefix is not None:
            # Derive slide title: use prefix + filename (without extension)
            base_name = filename.rsplit(".", 1)[0] if "." in filename else filename
            title_text = f"{title_prefix} {base_name}" if title_prefix else base_name

            txBox = slide.shapes.add_textbox(
                0, image_area_h, slide_w, title_height
            )
            tf = txBox.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = title_text
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
